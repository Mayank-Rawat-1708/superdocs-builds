"""
@file: backend/agents/nodes/compare.py
@description: Quarter-over-quarter comparison. Parses last quarter's digest, extracts
    its themes and volumes, matches them to this quarter's themes by embedding
    similarity (so a renamed theme is still recognised), and computes growth. When no
    prior digest is supplied, the node says so explicitly and every theme is reported as
    "no comparison available" rather than silently implied to be new.
@flow: run() -> if no prior digest, mark comparison unavailable and skip -> else read
    the file (DOCX/TXT/MD) -> LLM extracts prior themes and volumes -> embed both sides
    -> greedy one-to-one match above threshold -> write prior_quarter_count, growth_rate
    and volume_trend onto each Theme -> record disappeared themes as RESOLVED findings.
@dependencies:
    - backend.services.vector_store.match_prior_themes: greedy similarity matching
    - python-docx (optional): reading a .docx prior digest
"""

from __future__ import annotations

import logging
import uuid
from pathlib import Path

from sqlalchemy import select

from backend.agents.nodes.base import BaseNode, NodeSkip
from backend.agents.state import DigestState
from backend.db.database import session_scope
from backend.models import RunStatus, Theme, TrendDirection
from backend.services.groq_client import INJECTION_GUARD, GroqClient
from backend.services.heuristics import heuristic_parse_prior_digest
from backend.services.llm_gate import run_with_fallback
from backend.services.token_budget import fit_untrusted_excerpt
from backend.services.vector_store import get_embedder, match_prior_themes

logger = logging.getLogger(__name__)

# Below this cosine similarity two themes are considered different issues rather than
# the same issue renamed. Deliberately permissive: a false match is visible in the
# digest and correctable at the gate, a missed match silently inflates "new themes".
MATCH_THRESHOLD = 0.55

# Percentage change beyond which a theme is called GREW or SHRANK rather than STABLE.
STABLE_BAND = 0.15

_SYSTEM = """\
You read a previous quarter's Voice-of-Customer digest and extract its themes.

Return ONLY:
{"themes": [{"name": "<theme name as written>", "volume": <int or null>,
             "description": "<one sentence or empty string>"}]}

Rules:
- Copy theme names as they appear. Do not rewrite or normalise them.
- volume is the conversation count for that theme if the document states one; null if
  it does not. Never estimate a number the document does not contain.
- If the document contains no identifiable themes, return {"themes": []}.
"""


def _read_text(path: Path) -> str:
    """Extract plain text from the prior digest. Supports .docx, .txt, .md, .html."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        try:
            import docx  # python-docx
        except ImportError as exc:
            raise RuntimeError(
                "Reading a .docx prior digest needs python-docx (pip install python-docx)"
            ) from exc
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if suffix in {".txt", ".md", ".html", ".htm"}:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported prior-digest format {suffix!r}")


class CompareNode(BaseNode):
    stage = "compare"
    running_status = RunStatus.COMPARING
    # A missing or unreadable prior digest degrades the digest, it does not invalidate
    # this quarter's analysis, so a failure here must not fail the run.
    fatal_on_error = False

    async def run(self, state: DigestState, run_id: uuid.UUID) -> DigestState:
        # Nothing to compare against comes in two forms, and this is the one that used to
        # cost 77 seconds: no themes on THIS side. The node would still read the prior
        # digest, still send it to the model, still burn three retries and their backoff
        # on a request that had nowhere to put its answer. Comparison is a join between
        # two sets of themes — with one side empty the result is empty whatever the other
        # side says, so establish that before spending anything.
        async with session_scope() as session:
            current_theme_count = len(
                list(
                    (
                        await session.execute(
                            select(Theme.id).where(Theme.run_id == run_id)
                        )
                    ).scalars()
                )
            )
        if current_theme_count == 0:
            raise NodeSkip(
                "No themes in this quarter to compare against a prior digest",
                {
                    "comparison_available": False,
                    "prior_themes_found": 0,
                    "comparison_note": (
                        "This quarter produced no themes, so there is nothing to "
                        "compare against the prior digest."
                    ),
                },
            )

        prior_path_str = state.get("last_digest_path")
        if not prior_path_str:
            await self.log_decision(
                run_id,
                "NO_COMPARISON",
                "No prior-quarter digest supplied; QoQ comparison omitted rather than "
                "inferred",
            )
            raise NodeSkip(
                "No prior digest supplied",
                {
                    "comparison_available": False,
                    "prior_themes_found": 0,
                    "comparison_note": (
                        "No prior-quarter digest was supplied, so no quarter-over-quarter "
                        "comparison is possible. Themes are not marked as new or growing."
                    ),
                },
            )

        prior_path = Path(prior_path_str)
        if not prior_path.is_file():
            raise NodeSkip(
                f"Prior digest not found at {prior_path}",
                {
                    "comparison_available": False,
                    "prior_themes_found": 0,
                    "comparison_note": (
                        f"Prior digest {prior_path.name} could not be read, so no "
                        f"comparison is included."
                    ),
                },
            )

        raw_text = _read_text(prior_path)
        if not raw_text.strip():
            raise NodeSkip(
                f"{prior_path.name} contained no readable text",
                {
                    "comparison_available": False,
                    "prior_themes_found": 0,
                    "comparison_note": f"{prior_path.name} contained no readable text.",
                },
            )

        # A prior digest is as long as it is, and a fixed 20,000-character slice is not
        # a budget — 20,000 characters is roughly 6,000 tokens, which on an 8,000-token
        # ceiling leaves no room for the answer. Size the slice to what actually fits and
        # say how much was read, rather than sending a request that cannot be answered.
        answer_budget = 2048
        excerpt, excerpt_note = fit_untrusted_excerpt(
            raw_text,
            fixed_prompt=f"{INJECTION_GUARD}\n\n{_SYSTEM}"
            + "Extract the themes from this prior-quarter digest.",
            answer_tokens=answer_budget,
        )
        if excerpt_note:
            logger.info("Prior digest %s: %s", prior_path.name, excerpt_note)

        async def _llm():
            groq = GroqClient()
            try:
                payload, usage, _ = await groq.complete_json(
                    _SYSTEM,
                    "Extract the themes from this prior-quarter digest.",
                    untrusted_content=excerpt,
                    max_tokens=answer_budget,
                    min_completion_tokens=512,
                )
                return payload, usage
            finally:
                await groq.aclose()

        # Reading the prior digest has a real pattern-matching equivalent, so a request
        # that cannot fit degrades to it with disclosure rather than failing. There is no
        # batch here to split.
        gated = await run_with_fallback(
            "compare", _llm, lambda: heuristic_parse_prior_digest(raw_text),
            degrade_on_budget_error=True,
        )
        payload = gated.data
        self.usage.merge(gated.usage)
        if gated.degraded:
            await self.log_decision(
                run_id, "DEGRADED",
                f"The prior digest was parsed by pattern matching, not a language model "
                f"({gated.reason}). Themes stated in prose may have been missed.",
                {"caveats": gated.caveats},
            )
            state["degraded_stages"] = sorted(
                set(state.get("degraded_stages") or []) | {"compare"}
            )
            state["degraded_caveats"] = list(
                dict.fromkeys((state.get("degraded_caveats") or []) + gated.caveats)
            )

        prior_raw = payload.get("themes", []) if isinstance(payload, dict) else []
        prior_entries: list[tuple[str, str, int | None]] = []
        for item in prior_raw:
            name = str(item.get("name", "")).strip()
            if not name:
                continue
            volume = item.get("volume")
            try:
                volume_int = int(volume) if volume is not None else None
            except (TypeError, ValueError):
                volume_int = None
            prior_entries.append((name, str(item.get("description", "")), volume_int))

        if not prior_entries:
            raise NodeSkip(
                f"No themes could be extracted from {prior_path.name}",
                {
                    "comparison_available": False,
                    "prior_themes_found": 0,
                    "comparison_note": (
                        f"No themes could be identified in {prior_path.name}, so no "
                        f"comparison is included."
                    ),
                },
            )

        embedder = get_embedder()
        prior_vectors = embedder.embed(
            [f"{name}. {desc}" for name, desc, _ in prior_entries]
        )
        prior_for_match = [
            (name, vec, count if count is not None else 0)
            for (name, _desc, count), vec in zip(prior_entries, prior_vectors)
        ]

        async with session_scope() as session:
            themes = list(
                (
                    await session.execute(select(Theme).where(Theme.run_id == run_id))
                ).scalars()
            )
            current = [
                (t.name, list(t.embedding) if t.embedding is not None else [])
                for t in themes
            ]
            matches = match_prior_themes(
                current, prior_for_match, threshold=MATCH_THRESHOLD
            )

            new_count = 0
            for theme in themes:
                match = matches.get(theme.name)
                if match is None:
                    theme.volume_trend = TrendDirection.NEW
                    theme.prior_quarter_count = None
                    theme.growth_rate = None
                    new_count += 1
                    continue

                prior_name, prior_count, similarity = match
                theme.prior_theme_name = prior_name
                theme.match_similarity = round(similarity, 4)
                theme.prior_quarter_count = prior_count

                if prior_count <= 0:
                    # The prior digest named the theme but gave no number. We can say it
                    # existed, not how it moved — so we do not fabricate a growth rate.
                    theme.growth_rate = None
                    theme.volume_trend = TrendDirection.STABLE
                    theme.confidence_note = (
                        (theme.confidence_note or "")
                        + f" Prior digest named '{prior_name}' but stated no volume, so "
                        f"the change cannot be quantified."
                    ).strip()
                    continue

                growth = (theme.volume_count - prior_count) / prior_count
                theme.growth_rate = round(growth, 4)
                if growth > STABLE_BAND:
                    theme.volume_trend = TrendDirection.GREW
                elif growth < -STABLE_BAND:
                    theme.volume_trend = TrendDirection.SHRANK
                else:
                    theme.volume_trend = TrendDirection.STABLE

            # Prior themes with no current match: report as resolved/disappeared rather
            # than dropping them, because their absence is itself a finding.
            matched_prior = {m[0] for m in matches.values()}
            disappeared = [
                {"name": name, "prior_volume": count, "description": desc}
                for name, desc, count in prior_entries
                if name not in matched_prior
            ]
            await session.flush()

        note = (
            f"Compared against {len(prior_entries)} themes from {prior_path.name}. "
            f"{len(matches)} matched, {new_count} new, {len(disappeared)} no longer present."
        )
        await self.log_decision(
            run_id,
            "COMPARED",
            note,
            {
                "prior_themes": len(prior_entries),
                "matched": len(matches),
                "new": new_count,
                "disappeared": [d["name"] for d in disappeared],
            },
        )

        state["comparison_available"] = True
        state["prior_themes_found"] = len(prior_entries)
        state["comparison_note"] = note
        state["disappeared_themes"] = disappeared
        return state
