"""
@file: backend/services/local_render.py
@description: Renders the finished digest locally, without SuperDocs. Used when the
    SuperDocs key is missing, rejected, out of quota, or the service is unreachable —
    so a run that has completed every analysis stage still produces the document rather
    than dying at the last step with all the work stranded in the database.
@flow: superdocs_node catches an unavailability error -> calls render_digest() with the
    same approved draft sections it would have sent as edit instructions -> a DOCX (or
    Markdown fallback) is written to disk and recorded on the run exactly as a SuperDocs
    export would be.
@dependencies:
    - python-docx (optional): DOCX output; Markdown is written when it is absent

Note on fidelity: this is a plain renderer, not a substitute for SuperDocs' editing. It
lays out the sections we generated in order with basic styling. The digest states which
path produced it, so nobody mistakes a locally-rendered document for one that went
through the review-and-edit pipeline.
"""

from __future__ import annotations

import html
import logging
import re
from datetime import date
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

# Human-readable headings for each draft section key, in document order.
SECTION_TITLES: list[tuple[str, str]] = [
    ("executive_summary", "Executive Summary"),
    ("theme_table", "Top Themes This Quarter"),
    ("chart_volume", "Volume by Theme"),
    ("fastest_growing", "Fastest Growing Issues"),
    ("chart_growth", "Growth Chart"),
    ("what_changed", "What Changed"),
    ("chart_comparison", "Quarter-over-Quarter Comparison"),
    ("methodology", "Methodology"),
]

# Draft sections are phrased as instructions to an editor ("Replace the X section body
# with exactly:"). Rendering locally means stripping that framing to leave the content.
_INSTRUCTION_PREFIXES = re.compile(
    r"^(?:"
    r"Replace the [^\n:]*? with exactly this text, unchanged:"
    r"|Replace the [^\n:]*? with exactly:"
    r"|In the [^\n:]*?, insert (?:a table with the columns[^\n:]*?|this chart markup exactly as given[^\n:]*?):"
    r"|In the [^\n:]*?, immediately after the table, insert this chart markup exactly as given\."
    r"[^\n:]*?:"
    r"|Add a subsection under [^\n:]*? titled"
    r"|Set the document title to"
    r")\s*",
    re.IGNORECASE | re.DOTALL,
)

_SVG_BLOCK = re.compile(r"<svg\b.*?</svg>", re.DOTALL | re.IGNORECASE)
_PRE_BLOCK = re.compile(r"<pre>(.*?)</pre>", re.DOTALL | re.IGNORECASE)


def strip_instruction_framing(text: str) -> str:
    """Remove the 'tell the editor to do X' wrapper, leaving the content itself."""
    cleaned = _INSTRUCTION_PREFIXES.sub("", text.strip(), count=1)
    cleaned = re.sub(
        r"\n*Do not add commentary, causes, or recommendations beyond this text\.?\s*$",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"\n*Use exactly these rows and do not add, reorder, or invent any:?\s*",
        "\n",
        cleaned,
        flags=re.IGNORECASE,
    )
    return cleaned.strip()


def extract_chart_text(section: str) -> str | None:
    """Pull the block-character chart out of a chart section.

    The text form is used rather than the SVG because python-docx cannot embed SVG, and
    a chart the reader can actually see beats one that renders as a broken placeholder.
    """
    match = _PRE_BLOCK.search(section)
    return html.unescape(match.group(1)).strip() if match else None


# A per-theme deep dive is keyed theme_1, theme_2, ... The theme TABLE is keyed
# theme_table, which also starts with "theme_" — matching on the prefix alone rendered
# the table a second time under Theme Deep Dives, with its markdown header row promoted
# to a heading. Match the numeric form only.
_DEEP_DIVE_KEY = re.compile(r"^theme_(\d+)$")

# The instruction that introduces a deep-dive section, e.g.
#   Add a subsection under 'Theme Deep Dives' titled 'Export fails — 4 conversations
#   (New this quarter)'. Its body must contain, in order:
# Captures the title so it can become a real heading rather than being left inline.
_THEME_INSTRUCTION = re.compile(
    r"Add a subsection under[^']*'[^']*'\s*titled\s*'(?P<title>.+?)'\.\s*"
    r"Its body must contain,? in order:\s*",
    re.IGNORECASE | re.DOTALL,
)


def _deep_dive_keys(sections: dict[str, str]) -> list[str]:
    """Per-theme section keys in order. Excludes theme_table."""
    keyed = [(m.group(1), k) for k in sections if (m := _DEEP_DIVE_KEY.match(k))]
    return [k for _, k in sorted(keyed, key=lambda pair: int(pair[0]))]


def split_theme_section(section: str) -> tuple[str, str]:
    """Split a deep-dive instruction into (heading, body).

    Without this the whole instruction was emitted as one heading, so documents carried
    headings reading "Export fails — 4 conversations (New this quarter)'. Its body must
    contain, in order:" — the editor instruction leaking into the reader's document.
    """
    match = _THEME_INSTRUCTION.search(section)
    if match:
        title = match.group("title").strip()
        body = section[match.end():]
    else:
        # Unrecognised shape: fall back to generic stripping and use the first line as
        # the heading, rather than losing the content entirely.
        body = strip_instruction_framing(section)
        lines = [ln for ln in body.split("\n") if ln.strip()]
        title = lines[0].strip("'\" ") if lines else "Theme"
        body = "\n".join(lines[1:])

    body = re.sub(
        r"\n*Do not add commentary, causes, or recommendations beyond this text\.?\s*$",
        "", body, flags=re.IGNORECASE,
    )
    return title, body.strip()


def _is_table_block(text: str) -> bool:
    lines = [ln for ln in text.split("\n") if ln.strip()]
    return len(lines) >= 2 and sum(1 for ln in lines if ln.strip().startswith("|")) >= 2


def _parse_table(text: str) -> list[list[str]]:
    rows = []
    for line in text.split("\n"):
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        # Skip the markdown separator row.
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def render_digest_markdown(
    sections: dict[str, str],
    *,
    quarter_label: str,
    conversation_count: int,
    reason: str,
) -> str:
    """Markdown digest. Always available — no optional dependency."""
    parts = [
        f"# {quarter_label} Voice-of-Customer Digest",
        "",
        f"*Generated {date.today().isoformat()} from {conversation_count} support conversations.*",
        "",
        f"> **Rendered locally.** {reason} This document was produced directly from the "
        f"analysis rather than through the SuperDocs editing pipeline, so it has not been "
        f"through document-level AI editing. All content, figures and citations are "
        f"unchanged.",
        "",
    ]

    for key, title in SECTION_TITLES:
        if key not in sections:
            continue
        body = strip_instruction_framing(sections[key])
        parts.append(f"## {title}")
        parts.append("")
        chart = extract_chart_text(sections[key])
        if chart:
            parts += ["```", chart, "```", ""]
        else:
            parts += [_SVG_BLOCK.sub("", body).strip(), ""]

    if (theme_keys := _deep_dive_keys(sections)):
        parts += ["## Theme Deep Dives", ""]
        for key in theme_keys:
            title, body = split_theme_section(sections[key])
            parts += [f"### {title}", "", body, ""]

    return "\n".join(parts).rstrip() + "\n"


def render_digest_docx(
    sections: dict[str, str],
    out_path: Path,
    *,
    quarter_label: str,
    conversation_count: int,
    reason: str,
) -> Path:
    """DOCX digest. Falls back to Markdown when python-docx is unavailable."""
    try:
        import docx
        from docx.shared import Pt, RGBColor
    except ImportError:
        logger.warning("python-docx unavailable; writing Markdown instead of DOCX")
        md_path = out_path.with_suffix(".md")
        md_path.write_text(
            render_digest_markdown(
                sections,
                quarter_label=quarter_label,
                conversation_count=conversation_count,
                reason=reason,
            ),
            encoding="utf-8",
        )
        return md_path

    document = docx.Document()
    document.add_heading(f"{quarter_label} Voice-of-Customer Digest", level=0)
    document.add_paragraph(
        f"Generated {date.today().isoformat()} from {conversation_count} support conversations."
    )

    # Provenance banner. A reader must be able to tell which pipeline produced this.
    banner = document.add_paragraph()
    run = banner.add_run(
        f"Rendered locally. {reason} This document was produced directly from the "
        f"analysis rather than through the SuperDocs editing pipeline. All content, "
        f"figures and citations are unchanged."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    for key, title in SECTION_TITLES:
        if key not in sections:
            continue
        document.add_heading(title, level=1)
        raw = sections[key]
        body = strip_instruction_framing(raw)

        chart = extract_chart_text(raw)
        if chart:
            # Monospace so the block-character bars line up.
            para = document.add_paragraph()
            chart_run = para.add_run(chart)
            chart_run.font.name = "Courier New"
            chart_run.font.size = Pt(7.5)
            continue

        body = _SVG_BLOCK.sub("", body).strip()
        if _is_table_block(body):
            rows = _parse_table(body)
            if rows:
                table = document.add_table(rows=1, cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for i, cell_text in enumerate(rows[0]):
                    table.rows[0].cells[i].text = cell_text
                for row in rows[1:]:
                    cells = table.add_row().cells
                    for i, cell_text in enumerate(row[: len(cells)]):
                        cells[i].text = cell_text
            continue

        for paragraph in [p for p in body.split("\n") if p.strip()]:
            document.add_paragraph(paragraph.strip())

    if (theme_keys := _deep_dive_keys(sections)):
        document.add_heading("Theme Deep Dives", level=1)
        for key in theme_keys:
            title, body = split_theme_section(sections[key])
            document.add_heading(title, level=2)
            for line in [ln for ln in body.split("\n") if ln.strip()]:
                document.add_paragraph(line.strip())

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    return out_path


def render_digest(
    sections: dict[str, str],
    out_path: Path,
    *,
    quarter_label: str = "This quarter",
    conversation_count: int = 0,
    reason: str = "SuperDocs was unavailable.",
    fmt: str = "docx",
) -> dict[str, Any]:
    """Render the digest locally and report what was produced."""
    if fmt == "markdown":
        out_path = out_path.with_suffix(".md")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(
            render_digest_markdown(
                sections,
                quarter_label=quarter_label,
                conversation_count=conversation_count,
                reason=reason,
            ),
            encoding="utf-8",
        )
    else:
        out_path = render_digest_docx(
            sections,
            out_path,
            quarter_label=quarter_label,
            conversation_count=conversation_count,
            reason=reason,
        )

    logger.info("Locally rendered digest -> %s (%d bytes)", out_path, out_path.stat().st_size)
    return {
        "path": str(out_path),
        "format": out_path.suffix.lstrip("."),
        "bytes": out_path.stat().st_size,
        "sections_rendered": len(
            [k for k in sections if k in dict(SECTION_TITLES) or k.startswith("theme_")]
        ),
        "rendered_locally": True,
        "reason": reason,
    }